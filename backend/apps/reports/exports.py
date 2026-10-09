"""Exports Excel et PDF d'un rapport (structure commune produite par `services`)."""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.core.pdf import render_pdf
from apps.settings_app.models import LibrarySettings

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMATS = {"money": '#,##0" FCFA"', "int": "0", "percent": '0.0" %"', "decimal": "0.00"}
HEADER = {"font": Font(bold=True, color="FFFFFF"), "fill": PatternFill("solid", fgColor="1F5FAD")}


def filename(data: dict, extension: str) -> str:
    period = data["period"]
    return f"{data['key']}-{period['date_from']}-{period['date_to']}.{extension}"


def _sheet_title(title: str, used: set) -> str:
    base = "".join(c for c in title if c not in '[]:*?/\\')[:28] or "Feuille"
    name, n = base, 2
    while name in used:
        name, n = f"{base[:25]} {n}", n + 1
    used.add(name)
    return name


def to_xlsx(data: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Synthèse"
    period = data["period"]
    summary.append([f"{data['title']} — du {period['date_from']} au {period['date_to']} ({period['depot_name']})"])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append([])
    for item in data["summary"]:
        summary.append([item["label"], item["value"]])
        cell = summary.cell(row=summary.max_row, column=2)
        cell.number_format = FORMATS.get(item["kind"], "General")
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 18
    used = {"Synthèse"}
    for block in data["tables"]:
        sheet = workbook.create_sheet(_sheet_title(block["title"], used))
        sheet.append([c["label"] for c in block["columns"]])
        for cell in sheet[1]:
            cell.font, cell.fill = HEADER["font"], HEADER["fill"]
        for row in block["rows"]:
            sheet.append([row.get(c["key"]) for c in block["columns"]])
        for index, column in enumerate(block["columns"], start=1):
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = 40 if index == 1 else 16
            if column["kind"] in FORMATS:
                for cell in sheet[letter][1:]:
                    cell.number_format = FORMATS[column["kind"]]
        sheet.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_pdf(data: dict) -> bytes:
    """PDF A4 paysage (lève `PdfUnavailable` si WeasyPrint n'est pas utilisable)."""
    tables = [{**block, "cells": [[(row.get(c["key"]), c["kind"]) for c in block["columns"]] for row in block["rows"]]}
              for block in data["tables"]]
    return render_pdf("reports/report.html", {"report": data, "tables": tables, "library": LibrarySettings.load(),
                                              "generated_at": timezone.now()})
