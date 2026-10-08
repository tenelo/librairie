from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import extend_schema
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework import mixins, serializers, status, viewsets
from rest_framework.decorators import action
from rest_framework.response import Response

from apps.accounts.models import User
from apps.audit.models import AuditLog
from apps.audit.services import log_event
from apps.catalog.models import Article, Category
from apps.core.numbering import next_document_number
from apps.core.pdf import PdfUnavailable, render_pdf
from apps.core.permissions import HasAppPermission, IsStaffMember
from apps.settings_app.models import LibrarySettings
from apps.stock.models import Depot, Location

from . import services
from .models import InventoryAssignment, InventoryCount, InventorySession

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class AssignmentSerializer(serializers.ModelSerializer):
    user = serializers.SlugRelatedField(slug_field="username", queryset=User.objects.filter(is_active=True))
    user_name = serializers.CharField(source="user.get_full_name", read_only=True)
    categories = serializers.PrimaryKeyRelatedField(queryset=Category.objects.all(), many=True, required=False)

    class Meta:
        model = InventoryAssignment
        fields = ["id", "user", "user_name", "categories", "zone"]
        read_only_fields = ["id"]


class InventorySessionSerializer(serializers.ModelSerializer):
    depot = serializers.PrimaryKeyRelatedField(queryset=Depot.objects.filter(is_active=True))
    categories = serializers.PrimaryKeyRelatedField(queryset=Category.objects.all(), many=True, required=False)
    locations = serializers.PrimaryKeyRelatedField(queryset=Location.objects.all(), many=True, required=False)
    articles = serializers.PrimaryKeyRelatedField(queryset=Article.objects.all(), many=True, required=False)
    assignments = AssignmentSerializer(many=True, required=False)
    kind_label = serializers.CharField(source="get_kind_display", read_only=True)
    status_label = serializers.CharField(source="get_status_display", read_only=True)
    depot_name = serializers.CharField(source="depot.name", read_only=True)
    created_by_name = serializers.CharField(source="created_by.get_full_name", read_only=True)
    validated_by_name = serializers.CharField(source="validated_by.get_full_name", read_only=True, default=None)
    summary = serializers.SerializerMethodField()

    class Meta:
        model = InventorySession
        fields = ["id", "number", "name", "kind", "kind_label", "status", "status_label", "depot", "depot_name",
                  "categories", "locations", "articles", "counting_mode", "double_count_threshold", "assignments",
                  "frozen_at", "validated_at", "closed_at", "created_by_name", "validated_by_name", "notes",
                  "summary", "created_at"]
        read_only_fields = ["number", "status", "frozen_at", "validated_at", "closed_at", "created_at"]

    def get_summary(self, obj) -> dict:
        return services.summary(obj) if obj.status != InventorySession.Status.DRAFT else {}

    def validate(self, attrs):
        if self.instance is not None and self.instance.status != InventorySession.Status.DRAFT:
            changed = set(attrs) - {"assignments", "notes", "name"}
            if changed:
                raise serializers.ValidationError("Le périmètre ne peut plus être modifié après le démarrage.")
        kind = attrs.get("kind", getattr(self.instance, "kind", None))
        if kind == InventorySession.Kind.CYCLE and not attrs.get("articles", self.instance and self.instance.articles.all()):
            raise serializers.ValidationError({"articles": "Sélectionnez les articles de l'inventaire tournant."})
        return attrs

    def _save_assignments(self, session, assignments):
        session.assignments.all().delete()
        for item in assignments:
            categories = item.pop("categories", [])
            assignment = InventoryAssignment.objects.create(session=session, **item)
            assignment.categories.set(categories)

    @transaction.atomic
    def create(self, validated_data):
        assignments = validated_data.pop("assignments", [])
        m2m = {name: validated_data.pop(name, []) for name in ("categories", "locations", "articles")}
        session = InventorySession.objects.create(number=next_document_number("INV"),
                                                  created_by=self.context["request"].user, **validated_data)
        for name, values in m2m.items():
            getattr(session, name).set(values)
        self._save_assignments(session, assignments)
        return session

    @transaction.atomic
    def update(self, instance, validated_data):
        assignments = validated_data.pop("assignments", None)
        for name in ("categories", "locations", "articles"):
            if name in validated_data:
                getattr(instance, name).set(validated_data.pop(name))
        for name, value in validated_data.items():
            setattr(instance, name, value)
        instance.save()
        if assignments is not None:
            self._save_assignments(instance, assignments)
        return instance


class InventoryCountSerializer(serializers.ModelSerializer):
    article_code = serializers.CharField(source="article.code", read_only=True)
    article_name = serializers.CharField(source="article.name", read_only=True)
    user_name = serializers.CharField(source="user.get_full_name", read_only=True)

    class Meta:
        model = InventoryCount
        fields = ["id", "article", "article_code", "article_name", "quantity", "count_round", "location",
                  "user_name", "counted_at"]


class ValidateSerializer(serializers.Serializer):
    uncounted_as_zero = serializers.BooleanField(default=False)


class InventorySessionViewSet(mixins.ListModelMixin, mixins.RetrieveModelMixin, mixins.CreateModelMixin,
                              mixins.UpdateModelMixin, viewsets.GenericViewSet):
    serializer_class = InventorySessionSerializer
    permission_classes = [IsStaffMember, HasAppPermission]
    required_permissions = {
        "list": None, "retrieve": None, "report": "inventory.view", "counts": "inventory.view",
        "export": "inventory.view", "pdf": "inventory.view",
        "validate": "inventory.validate", "close": "inventory.validate", "*": "inventory.create",
    }
    filterset_fields = ["status", "kind", "depot"]
    search_fields = ["number", "name"]

    def get_queryset(self):
        queryset = InventorySession.objects.select_related("depot", "created_by", "validated_by").prefetch_related(
            "categories", "locations", "articles", "assignments__user", "assignments__categories")
        user = self.request.user
        if not (user.has_app_perm("inventory.view") or user.has_app_perm("inventory.create")):
            # Un inventoriste ne voit que les inventaires en cours.
            queryset = queryset.filter(status__in=InventorySession.COUNTABLE)
        return queryset

    def _transition(self, func, label, **kwargs):
        session = func(self.get_object(), self.request.user, **kwargs)
        log_event(AuditLog.Action.UPDATE, request=self.request, entity="InventorySession", entity_id=session.pk,
                  description=f"{label} de l'inventaire {session.number}")
        return Response(self.get_serializer(self.get_queryset().get(pk=session.pk)).data)

    @extend_schema(request=None, responses=InventorySessionSerializer)
    @action(detail=True, methods=["post"])
    def start(self, request, pk=None):
        return self._transition(services.start, "Démarrage")

    @extend_schema(request=None, responses=InventorySessionSerializer)
    @action(detail=True, methods=["post"])
    def review(self, request, pk=None):
        return self._transition(services.to_review, "Passage en revue")

    @extend_schema(request=ValidateSerializer, responses=InventorySessionSerializer)
    @action(detail=True, methods=["post"])
    def validate(self, request, pk=None):
        params = ValidateSerializer(data=request.data)
        params.is_valid(raise_exception=True)
        session = services.validate(self.get_object(), request.user, request=request,
                                    uncounted_as_zero=params.validated_data["uncounted_as_zero"])
        return Response(self.get_serializer(self.get_queryset().get(pk=session.pk)).data)

    @extend_schema(request=None, responses=InventorySessionSerializer)
    @action(detail=True, methods=["post"])
    def close(self, request, pk=None):
        session = self.get_object()
        if session.status != InventorySession.Status.VALIDATED:
            return Response({"detail": "Seul un inventaire validé peut être clôturé."}, status=400)
        session.status = InventorySession.Status.CLOSED
        session.closed_at = timezone.now()
        session.save(update_fields=["status", "closed_at"])
        return Response(self.get_serializer(session).data)

    @extend_schema(request=None, responses=InventorySessionSerializer)
    @action(detail=True, methods=["post"])
    def cancel(self, request, pk=None):
        session = self.get_object()
        if session.status in (InventorySession.Status.VALIDATED, InventorySession.Status.CLOSED):
            return Response({"detail": "Un inventaire validé ne peut pas être annulé."}, status=400)
        session.status = InventorySession.Status.CANCELLED
        session.save(update_fields=["status"])
        return Response(self.get_serializer(session).data)

    def _lines(self, session):
        if session.status in (InventorySession.Status.VALIDATED, InventorySession.Status.CLOSED):
            return [{"article": str(line.article_id), "code": line.article.code, "name": line.article.name,
                     "theoretical": line.theoretical, "movements_during": line.movements_during,
                     "expected": line.expected, "first_count": None, "second_count": None, "counted": line.counted,
                     "variance": line.variance, "unit_cost": line.unit_cost, "variance_value": line.variance_value,
                     "needs_recount": False} for line in session.lines.select_related("article")]
        return [line.as_dict() for line in services.variance_report(session)]

    @extend_schema(responses=OpenApiTypes.OBJECT)
    @action(detail=True, methods=["get"])
    def report(self, request, pk=None):
        """Rapport d'écarts : théorique, mouvements pendant l'inventaire, compté, écart, valeur."""
        session = self.get_object()
        return Response({"lines": self._lines(session), "summary": services.summary(session)})

    @extend_schema(responses=InventoryCountSerializer(many=True))
    @action(detail=True, methods=["get"])
    def counts(self, request, pk=None):
        counts = self.get_object().counts.select_related("article", "user")
        if request.query_params.get("mine") in ("1", "true"):
            counts = counts.filter(user=request.user)
        return Response(InventoryCountSerializer(counts, many=True).data)

    @extend_schema(responses={(200, XLSX): bytes})
    @action(detail=True, methods=["get"])
    def export(self, request, pk=None):
        session = self.get_object()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Écarts"
        headers = ["Code", "Article", "Théorique", "Mouvements pendant", "Attendu", "Compté", "Écart", "CMP",
                   "Valeur de l'écart"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F5FAD")
        for line in self._lines(session):
            sheet.append([line["code"], line["name"], line["theoretical"], line["movements_during"], line["expected"],
                          line["counted"], line["variance"], line["unit_cost"], line["variance_value"]])
        sheet.column_dimensions["B"].width = 40
        buffer = BytesIO()
        workbook.save(buffer)
        response = HttpResponse(buffer.getvalue(), content_type=XLSX)
        response["Content-Disposition"] = f'attachment; filename="ecarts-{session.number}.xlsx"'
        return response

    @extend_schema(responses={(200, "application/pdf"): bytes})
    @action(detail=True, methods=["get"])
    def pdf(self, request, pk=None):
        session = self.get_object()
        lines = self._lines(session)
        try:
            content = render_pdf("inventory/variance_report.html", {
                "session": session, "lines": lines, "summary": services.summary(session),
                "library": LibrarySettings.load(), "generated_at": timezone.now()})
        except PdfUnavailable as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_503_SERVICE_UNAVAILABLE)
        response = HttpResponse(content, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="ecarts-{session.number}.pdf"'
        return response
