from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.catalog.models import Article, Barcode, Category, Kit, PriceHistory
from apps.customers.models import School

from .factories import make_article, make_kit

pytestmark = pytest.mark.django_db

ARTICLE_HEADERS = ["Code", "Désignation", "Type", "Catégorie", "Marque", "Unité", "Code-barres", "Prix d'achat",
                   "Prix détail", "Prix gros", "Seuil gros", "TVA %", "Seuil alerte", "Stock max", "Actif",
                   "En ligne"]


def xlsx(headers, rows, name="import.xlsx"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue(),
                              content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def test_import_articles_creates_and_updates(auth_client, manager_user):
    existing = make_article("Ancien nom", retail=300, code="ART-90001")
    upload = xlsx(ARTICLE_HEADERS, [
        ["", "Cahier 100 pages", "SIMPLE", "Cahiers > Grands cahiers", "Clairefontaine", "Unité", "4006381333931",
         250, 400, 350, 10, None, 20, 500, "oui", "oui"],
        ["ART-90001", "Nouveau nom", "SIMPLE", "Écriture", "", "", "", None, "350", None, None, None, 5, None, "oui",
         "non"],
        ["", "Photocopie A4", "SERVICE", "Services", "", "Page", "", 0, 25, None, None, None, 0, None, "oui", "non"],
    ])
    response = auth_client(manager_user).post("/api/catalog/import/articles/", {"file": upload}, format="multipart")
    assert response.status_code == 200, response.content
    result = response.json()
    assert result == {"created": 2, "updated": 1, "errors": [], "dry_run": False, "applied": True}

    cahier = Article.objects.get(name="Cahier 100 pages")
    assert cahier.category.path == "Cahiers › Grands cahiers"
    assert cahier.brand.name == "Clairefontaine" and cahier.wholesale_min_qty == 10
    assert Barcode.objects.filter(article=cahier, code="4006381333931", kind="EAN13").exists()
    existing.refresh_from_db()
    assert existing.name == "Nouveau nom" and existing.retail_price == 350
    assert PriceHistory.objects.filter(article=existing, reason="Import Excel").exists()
    assert Article.objects.get(name="Photocopie A4").barcodes.count() == 1


def test_import_dry_run_writes_nothing(auth_client, manager_user):
    upload = xlsx(ARTICLE_HEADERS, [["", "Règle 30 cm", "SIMPLE", "Écriture", "", "", "", 150, 300]])
    result = auth_client(manager_user).post("/api/catalog/import/articles/", {"file": upload, "dry_run": True},
                                            format="multipart").json()
    assert result["created"] == 1 and result["applied"] is False
    assert not Article.objects.filter(name="Règle 30 cm").exists()


def test_import_errors_are_reported_and_nothing_is_saved(auth_client, manager_user):
    upload = xlsx(ARTICLE_HEADERS, [
        ["", "Bon article", "SIMPLE", "", "", "", "", 100, 200],
        ["", "", "SIMPLE", "", "", "", "", 100, 200],
        ["", "Prix faux", "SIMPLE", "", "", "", "", 100, "abc"],
        ["", "Type faux", "KIT", "", "", "", "", 100, 200],
        ["", "Unité fausse", "SIMPLE", "", "", "Tonneau", "", 100, 200],
    ])
    result = auth_client(manager_user).post("/api/catalog/import/articles/", {"file": upload},
                                            format="multipart").json()
    assert [e["row"] for e in result["errors"]] == [3, 4, 5, 6]
    assert result["applied"] is False
    assert not Article.objects.filter(name="Bon article").exists()


def test_import_requires_permission(auth_client, cashier_user):
    upload = xlsx(ARTICLE_HEADERS, [["", "X", "SIMPLE"]])
    assert auth_client(cashier_user).post("/api/catalog/import/articles/", {"file": upload},
                                          format="multipart").status_code == 403


def test_export_articles_roundtrip(auth_client, manager_user):
    article = make_article("Stylo bleu", retail=150, category=Category.objects.get(name="Écriture"))
    Barcode.objects.create(article=article, code="4006381333931")
    response = auth_client(manager_user).get("/api/catalog/articles/export/")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:3] == ("Code", "Désignation", "Type")
    assert rows[1][1] == "Stylo bleu" and rows[1][6] == "4006381333931" and rows[1][8] == 150

    upload = SimpleUploadedFile("export.xlsx", response.content)
    result = auth_client(manager_user).post("/api/catalog/import/articles/", {"file": upload},
                                            format="multipart").json()
    assert result["errors"] == [] and result["updated"] == 1


KIT_HEADERS = ["Code kit", "Nom du kit", "Niveau", "Établissement", "Année scolaire", "Code article", "Quantité",
               "Facultatif"]


def test_import_kit_lists(auth_client, manager_user):
    cahier = make_article("Cahier", 400, code="ART-80001")
    stylo = make_article("Stylo", 150, code="ART-80002")
    School.objects.create(name="Lycée Moderne de Cocody", city="Abidjan")
    upload = xlsx(KIT_HEADERS, [
        ["", "Kit 6ème LMC", "6EME", "Lycée Moderne de Cocody", "2026-2027", "ART-80001", 6, "non"],
        ["", "Kit 6ème LMC", "6EME", "Lycée Moderne de Cocody", "2026-2027", "ART-80002", 4, "oui"],
        ["", "Kit CP1", "CP1", "", "", "ART-80001", 2, ""],
    ])
    result = auth_client(manager_user).post("/api/catalog/import/kits/", {"file": upload},
                                            format="multipart").json()
    assert result["errors"] == [] and result["created"] == 2
    kit = Kit.objects.get(article__name="Kit 6ème LMC")
    assert kit.school.name == "Lycée Moderne de Cocody" and kit.school_year.label == "2026-2027"
    assert kit.article.retail_price == 6 * 400 + 4 * 150
    assert kit.components.get(article=stylo).is_optional

    # Réimport : la composition est remplacée, pas dupliquée.
    upload = xlsx(KIT_HEADERS, [[kit.article.code, "Kit 6ème LMC", "6EME", "Lycée Moderne de Cocody", "2026-2027",
                                 "ART-80001", 8, "non"]])
    result = auth_client(manager_user).post("/api/catalog/import/kits/", {"file": upload},
                                            format="multipart").json()
    assert result["updated"] == 1
    assert [(c.article_id, c.quantity) for c in kit.components.all()] == [(cahier.pk, 8)]


def test_import_kit_errors(auth_client, manager_user):
    upload = xlsx(KIT_HEADERS, [
        ["", "Kit", "9EME", "", "", "ART-00001", 1, ""],
        ["", "Kit", "CP1", "École inconnue", "", "ART-00001", 1, ""],
        ["", "Kit", "CP1", "", "", "INCONNU", 1, ""],
    ])
    result = auth_client(manager_user).post("/api/catalog/import/kits/", {"file": upload},
                                            format="multipart").json()
    assert len(result["errors"]) == 3 and Kit.objects.count() == 0


def test_export_kits_is_reimportable(auth_client, manager_user):
    make_kit([(make_article("Cahier", 400), 6)], name="Kit 6ème")
    response = auth_client(manager_user).get("/api/catalog/kits/export/")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Listes de fournitures", "Prix"]
    upload = SimpleUploadedFile("kits.xlsx", response.content)
    result = auth_client(manager_user).post("/api/catalog/import/kits/", {"file": upload},
                                            format="multipart").json()
    assert result["errors"] == [] and result["updated"] == 1


def test_templates_download(auth_client, manager_user):
    client = auth_client(manager_user)
    for kind in ("articles", "kits"):
        response = client.get(f"/api/catalog/templates/{kind}/")
        assert response.status_code == 200
        expected_rows = 2 if kind == "articles" else 3
        assert load_workbook(BytesIO(response.content)).active.max_row == expected_rows
