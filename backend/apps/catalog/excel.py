"""Import / export Excel du catalogue et des listes de fournitures (kits scolaires).

Les imports sont transactionnels : en cas d'erreur sur une ligne, rien n'est enregistré
et le rapport indique chaque ligne fautive. Le mode « simulation » (dry_run) valide le
fichier et renvoie le même rapport sans rien écrire.
"""
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.customers.models import School

from .barcodes import is_valid_ean13
from .models import Article, Barcode, Brand, Category, Kit, KitComponent, SchoolLevel, SchoolYear, Unit
from .services import (
    PRICE_FIELDS,
    add_internal_barcode,
    compute_kit_pricing,
    next_article_code,
    price_snapshot,
    record_price_changes,
    refresh_kit_price,
    refresh_kits_using,
)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ARTICLE_COLUMNS = [
    ("code", "Code"),
    ("name", "Désignation"),
    ("type", "Type"),
    ("category", "Catégorie"),
    ("brand", "Marque"),
    ("unit", "Unité"),
    ("barcode", "Code-barres"),
    ("avg_cost", "Prix d'achat"),
    ("retail_price", "Prix détail"),
    ("wholesale_price", "Prix gros"),
    ("wholesale_min_qty", "Seuil gros"),
    ("vat_rate", "TVA %"),
    ("alert_threshold", "Seuil alerte"),
    ("max_stock", "Stock max"),
    ("is_active", "Actif"),
    ("is_online", "En ligne"),
    ("description", "Description"),
]

KIT_COLUMNS = [
    ("kit_code", "Code kit"),
    ("kit_name", "Nom du kit"),
    ("level", "Niveau"),
    ("school", "Établissement"),
    ("school_year", "Année scolaire"),
    ("article_code", "Code article"),
    ("quantity", "Quantité"),
    ("optional", "Facultatif"),
]

TYPE_BY_LABEL = {
    "SIMPLE": Article.Type.SIMPLE, "ARTICLE SIMPLE": Article.Type.SIMPLE,
    "LOT": Article.Type.LOT, "SERVICE": Article.Type.SERVICE,
}
TRUE_VALUES = {"oui", "o", "yes", "y", "vrai", "true", "1", "x"}
FALSE_VALUES = {"non", "n", "no", "faux", "false", "0", ""}

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F5FAD")


class ImportReport:
    def __init__(self):
        self.created = 0
        self.updated = 0
        self.errors = []

    def error(self, row, message):
        self.errors.append({"row": row, "message": message})

    def as_dict(self, dry_run):
        return {"created": self.created, "updated": self.updated, "errors": self.errors, "dry_run": dry_run,
                "applied": not dry_run and not self.errors}


class DryRunRollback(Exception):
    """Annule la transaction d'un import en simulation ou en erreur."""


# --- Utilitaires -----------------------------------------------------------------------------


def _workbook_response_bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_sheet(sheet, headers, rows):
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append(row)
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[index - 1] or "")) for r in rows]) + 2
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width, 10), 60)
    sheet.freeze_panes = "A2"


def _rows(file, columns):
    """Lit la première feuille ; associe les colonnes par libellé d'en-tête (ordre libre)."""
    workbook = load_workbook(file, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(h).strip().lower() if h is not None else "" for h in next(iterator, [])]
    by_label = {label.lower(): key for key, label in columns}
    mapping = {i: by_label[h] for i, h in enumerate(headers) if h in by_label}
    missing = [label for key, label in columns[:2] if key not in mapping.values()]
    if missing:
        raise ValueError(f"Colonnes obligatoires absentes : {', '.join(missing)}.")
    for number, values in enumerate(iterator, start=2):
        if values is None or all(v in (None, "") for v in values):
            continue
        yield number, {mapping[i]: v for i, v in enumerate(values) if i in mapping}


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _int(value, field, allow_none=True):
    text = _text(value).replace(" ", "").replace(" ", "").replace("\xa0", "")
    if text == "":
        if allow_none:
            return None
        raise ValueError(f"{field} est obligatoire.")
    try:
        number = Decimal(text.replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"{field} : « {text} » n'est pas un nombre.") from exc
    if number < 0 or number != number.to_integral_value():
        raise ValueError(f"{field} doit être un entier positif (FCFA sans décimales).")
    return int(number)


def _decimal(value, field):
    text = _text(value).replace(",", ".")
    if text == "":
        return None
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{field} : « {text} » n'est pas un nombre.") from exc


def _bool(value, default):
    text = _text(value).lower()
    if text == "" and value is None:
        return default
    if text in TRUE_VALUES:
        return True
    if text in FALSE_VALUES:
        return False
    raise ValueError(f"Valeur « {text} » non reconnue (oui / non attendu).")


def category_from_path(path: str):
    """Retrouve ou crée une catégorie à partir d'un chemin « Parent > Enfant »."""
    parent = None
    for name in [p.strip() for p in path.replace("›", ">").split(">") if p.strip()]:
        category = Category.objects.filter(parent=parent, name__iexact=name).first()
        if category is None:
            category = Category.objects.create(parent=parent, name=name)
        parent = category
    return parent


# --- Articles ---------------------------------------------------------------------------------


def export_articles(queryset) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Articles"
    rows = []
    for article in queryset.select_related("category", "brand", "unit").prefetch_related("barcodes"):
        barcodes = [b.code for b in article.barcodes.all() if not b.is_deleted]
        rows.append([
            article.code, article.name, article.type, article.category.path if article.category else "",
            article.brand.name if article.brand else "", article.unit.label if article.unit else "",
            barcodes[0] if barcodes else "", article.avg_cost, article.retail_price, article.wholesale_price,
            article.wholesale_min_qty, float(article.vat_rate) if article.vat_rate is not None else None,
            article.alert_threshold, article.max_stock, "oui" if article.is_active else "non",
            "oui" if article.is_online else "non", article.description,
        ])
    _write_sheet(sheet, [label for _key, label in ARTICLE_COLUMNS], rows)
    return _workbook_response_bytes(workbook)


def articles_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Articles"
    _write_sheet(sheet, [label for _key, label in ARTICLE_COLUMNS], [[
        "", "Cahier 100 pages grands carreaux", "SIMPLE", "Cahiers > Grands cahiers", "Clairefontaine", "Unité",
        "", 250, 400, 350, 10, None, 20, 500, "oui", "oui", "Couverture souple",
    ]])
    return _workbook_response_bytes(workbook)


def import_articles(file, user, dry_run=False) -> dict:
    report = ImportReport()
    can_change_prices = user.has_app_perm("catalog.prices")
    try:
        with transaction.atomic():
            for number, row in _rows(file, ARTICLE_COLUMNS):
                try:
                    with transaction.atomic():
                        _import_article_row(row, user, can_change_prices, report)
                except ValueError as exc:
                    report.error(number, str(exc))
            if dry_run or report.errors:
                raise DryRunRollback
    except DryRunRollback:
        pass
    except ValueError as exc:
        report.error(1, str(exc))
    return report.as_dict(dry_run)


def _import_article_row(row, user, can_change_prices, report):
    name = _text(row.get("name"))
    if not name:
        raise ValueError("La désignation est obligatoire.")
    code = _text(row.get("code")).upper()
    type_label = _text(row.get("type")).upper() or "SIMPLE"
    article_type = TYPE_BY_LABEL.get(type_label)
    if article_type is None:
        raise ValueError(f"Type « {type_label} » invalide (SIMPLE, LOT ou SERVICE ; les kits s'importent à part).")

    article = Article.objects.filter(code__iexact=code).first() if code else None
    if article is not None and article.type == Article.Type.KIT:
        raise ValueError(f"{code} est un kit : utilisez l'import des listes de fournitures.")
    is_new = article is None
    if is_new:
        article = Article(code=code or next_article_code(article_type), type=article_type)
    elif article.type != article_type:
        raise ValueError(f"Le type de {code} ne peut pas être modifié ({article.type} → {article_type}).")

    before = price_snapshot(article) if not is_new else {f: None for f in PRICE_FIELDS}
    article.name = name
    if "description" in row:
        article.description = _text(row.get("description"))
    if _text(row.get("category")):
        article.category = category_from_path(_text(row["category"]))
    if _text(row.get("brand")):
        article.brand, _ = Brand.objects.get_or_create(name=_text(row["brand"]))
    if _text(row.get("unit")):
        label = _text(row["unit"])
        article.unit = Unit.objects.filter(label__iexact=label).first() or Unit.objects.filter(
            code__iexact=label).first()
        if article.unit is None:
            raise ValueError(f"Unité « {label} » inconnue.")

    prices = {
        "retail_price": _int(row.get("retail_price"), "Prix détail"),
        "wholesale_price": _int(row.get("wholesale_price"), "Prix gros"),
        "wholesale_min_qty": _int(row.get("wholesale_min_qty"), "Seuil gros"),
        "vat_rate": _decimal(row.get("vat_rate"), "TVA %"),
    }
    if is_new:
        prices["avg_cost"] = _int(row.get("avg_cost"), "Prix d'achat")
    for name_, value in prices.items():
        if name_ in ("wholesale_price", "wholesale_min_qty", "vat_rate"):
            if name_ in row:
                setattr(article, name_, value)
        elif value is not None:
            setattr(article, name_, value)
    changed = any(str(before.get(f)) != str(getattr(article, f)) for f in PRICE_FIELDS) and not (
        is_new and article.retail_price == 0 and article.avg_cost == 0
    )
    if changed and not can_change_prices:
        raise ValueError("Modification de prix refusée : permission « Modifier les prix » requise.")

    article.alert_threshold = _int(row.get("alert_threshold"), "Seuil alerte") or 0
    article.max_stock = _int(row.get("max_stock"), "Stock max")
    article.is_active = _bool(row.get("is_active"), True)
    article.is_online = _bool(row.get("is_online"), False)
    if article_type == Article.Type.LOT and not article.lot_base_article_id:
        raise ValueError("Les lots se créent depuis la fiche article (article contenu et nombre d'unités).")
    try:
        article.full_clean(exclude=["lot_base_article", "lot_units", "category", "brand", "unit"])
    except Exception as exc:
        messages = getattr(exc, "message_dict", {"": [str(exc)]})
        raise ValueError(" ".join(f"{k} : {', '.join(v)}" if k else ", ".join(v) for k, v in messages.items()))
    article.save()

    barcode = _text(row.get("barcode"))
    if barcode:
        owner = Barcode.objects.filter(code=barcode).exclude(article=article).select_related("article").first()
        if owner:
            raise ValueError(f"Le code-barres {barcode} appartient déjà à « {owner.article.name} ».")
        if not Barcode.objects.filter(code=barcode, article=article).exists():
            kind = Barcode.Kind.EAN13 if is_valid_ean13(barcode) else Barcode.Kind.OTHER
            Barcode.objects.create(article=article, code=barcode, kind=kind)
    elif is_new:
        add_internal_barcode(article)

    if not is_new:
        changes = record_price_changes(article, before, user=user, reason="Import Excel")
        if any(c.field == "retail_price" for c in changes):
            refresh_kits_using(article, user=user)
    if is_new:
        report.created += 1
    else:
        report.updated += 1


# --- Listes de fournitures (kits) ---------------------------------------------------------------


def export_kits(kits) -> bytes:
    """Une ligne par composant ; le fichier est réimportable tel quel."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Listes de fournitures"
    rows = []
    for kit in kits:
        for component in kit.components.all():
            rows.append([
                kit.article.code, kit.article.name, kit.level.code, kit.school.name if kit.school else "",
                kit.school_year.label if kit.school_year else "", component.article.code, component.quantity,
                "oui" if component.is_optional else "non",
            ])
    _write_sheet(sheet, [label for _key, label in KIT_COLUMNS], rows)

    summary = workbook.create_sheet("Prix")
    summary_rows = []
    for kit in kits:
        pricing = compute_kit_pricing(kit)
        summary_rows.append([kit.article.code, kit.article.name, kit.level.label,
                             kit.school.name if kit.school else "", pricing.components_total,
                             float(pricing.discount_percent), pricing.price, pricing.cost, pricing.margin])
    _write_sheet(summary, ["Code kit", "Kit", "Niveau", "Établissement", "Total composants", "Remise %",
                           "Prix", "Coût", "Marge"], summary_rows)
    return _workbook_response_bytes(workbook)


def kits_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Listes de fournitures"
    _write_sheet(sheet, [label for _key, label in KIT_COLUMNS], [
        ["", "Kit 6ème — Lycée Moderne", "6EME", "Lycée Moderne de Cocody", "2026-2027", "ART-00001", 6, "non"],
        ["", "Kit 6ème — Lycée Moderne", "6EME", "Lycée Moderne de Cocody", "2026-2027", "ART-00002", 2, "oui"],
    ])
    return _workbook_response_bytes(workbook)


def import_kits(file, user, dry_run=False) -> dict:
    """Regroupe les lignes par kit (code, ou nom + niveau + établissement + année) et remplace leur composition."""
    report = ImportReport()
    groups = {}
    try:
        for number, row in _rows(file, KIT_COLUMNS):
            try:
                key, header = _kit_header(row)
                article = Article.objects.filter(code__iexact=_text(row.get("article_code"))).first()
                if article is None:
                    raise ValueError(f"Article « {_text(row.get('article_code'))} » introuvable.")
                if article.type == Article.Type.KIT:
                    raise ValueError("Un kit ne peut pas contenir un autre kit.")
                quantity = _int(row.get("quantity"), "Quantité", allow_none=False)
                if quantity < 1:
                    raise ValueError("La quantité doit être d'au moins 1.")
                group = groups.setdefault(key, {"header": header, "lines": {}, "row": number})
                line = group["lines"].setdefault(article.pk, {"article": article, "quantity": 0, "optional": False})
                line["quantity"] += quantity
                line["optional"] = _bool(row.get("optional"), False)
            except ValueError as exc:
                report.error(number, str(exc))
    except ValueError as exc:
        report.error(1, str(exc))
        return report.as_dict(dry_run)

    try:
        with transaction.atomic():
            for group in groups.values():
                try:
                    _apply_kit_group(group, user, report)
                except ValueError as exc:
                    report.error(group["row"], str(exc))
            if dry_run or report.errors:
                raise DryRunRollback
    except DryRunRollback:
        pass
    return report.as_dict(dry_run)


def _kit_header(row):
    code = _text(row.get("kit_code")).upper()
    name = _text(row.get("kit_name"))
    level_code = _text(row.get("level")).upper()
    if not level_code:
        raise ValueError("Le niveau est obligatoire.")
    level = SchoolLevel.objects.filter(code__iexact=level_code).first()
    if level is None:
        raise ValueError(f"Niveau « {level_code} » inconnu.")
    school = None
    if _text(row.get("school")):
        school = School.objects.filter(name__iexact=_text(row["school"])).first()
        if school is None:
            raise ValueError(f"Établissement « {_text(row['school'])} » inconnu.")
    year = None
    if _text(row.get("school_year")):
        year = SchoolYear.objects.filter(label=_text(row["school_year"])).first()
        if year is None:
            raise ValueError(f"Année scolaire « {_text(row['school_year'])} » inconnue.")
    if not code and not name:
        raise ValueError("Indiquez le code ou le nom du kit.")
    key = code or (name.lower(), level.pk, school.pk if school else None, year.pk if year else None)
    return key, {"code": code, "name": name, "level": level, "school": school, "school_year": year}


def _apply_kit_group(group, user, report):
    header = group["header"]
    kit = None
    if header["code"]:
        kit = Kit.objects.filter(article__code__iexact=header["code"]).select_related("article").first()
    if kit is None and header["name"]:
        kit = Kit.objects.filter(
            article__name__iexact=header["name"], level=header["level"], school=header["school"],
            school_year=header["school_year"],
        ).select_related("article").first()
    if kit is None:
        if not header["name"]:
            raise ValueError(f"Kit {header['code']} introuvable : indiquez son nom pour le créer.")
        article = Article.objects.create(
            code=header["code"] or next_article_code(Article.Type.KIT), name=header["name"], type=Article.Type.KIT,
            category=Category.objects.filter(name__iexact="Kits scolaires", parent=None).first(),
        )
        add_internal_barcode(article)
        kit = Kit.objects.create(article=article, level=header["level"], school=header["school"],
                                 school_year=header["school_year"])
        report.created += 1
    else:
        if header["name"]:
            kit.article.name = header["name"]
            kit.article.save(update_fields=["name"])
        kit.level, kit.school, kit.school_year = header["level"], header["school"], header["school_year"]
        kit.save()
        report.updated += 1

    keep = set()
    for order, line in enumerate(group["lines"].values()):
        component = KitComponent.objects.filter(kit=kit, article=line["article"]).first()
        if component is None:
            component = KitComponent(kit=kit, article=line["article"])
        component.quantity = line["quantity"]
        component.is_optional = line["optional"]
        component.sort_order = order
        component.save()
        keep.add(component.pk)
    for component in kit.components.exclude(pk__in=keep):
        component.soft_delete()
    refresh_kit_price(kit, user=user)
